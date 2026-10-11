"""Update the Team Builder workbook: add any new events, then rewrite Meta_Cores / Core_Finder / Core_Data.
Keeps the user's tabs, settings and Excel-only dropdowns (re-injected after openpyxl saves)."""
import openpyxl, re, copy, os, zipfile, shutil, warnings, collections, itertools
from openpyxl.utils import get_column_letter as CL
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from tiers import tier, MARGIN
import analyse
warnings.simplefilter('ignore')
def cstyle(s,d):
    if s.has_style:
        d.font=copy.copy(s.font); d.fill=copy.copy(s.fill); d.border=copy.copy(s.border); d.alignment=copy.copy(s.alignment); d.number_format=s.number_format; d.protection=copy.copy(s.protection)
# ---------------- reference shifting (row insertion) ----------------
REF=re.compile(r"(?<![A-Za-z0-9_.])((?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!)?(\$?)([A-Z]{1,3})(\$?)(\d+)(?::(\$?)([A-Z]{1,3})(\$?)(\d+))?(?![A-Za-z0-9_(!])")
def shift(f,own,target,P,N):
    parts=f.split('"')
    for k in range(0,len(parts),2):
        def rep(m):
            sh=m.group(1); sheet=sh[:-1].strip("'") if sh else own
            if sheet!=target: return m.group(0)
            r1=int(m.group(5))
            if m.group(9) is None:
                if r1>=P: r1+=N
                return f"{sh or ''}{m.group(2)}{m.group(3)}{m.group(4)}{r1}"
            r2=int(m.group(9))
            if r1>=P: r1+=N; r2+=N
            elif r2>=P or (r2==P-1 and r2>r1): r2+=N
            return f"{sh or ''}{m.group(2)}{m.group(3)}{m.group(4)}{r1}:{m.group(6)}{m.group(7)}{m.group(8)}{r2}"
        parts[k]=REF.sub(rep,parts[k])
    return '"'.join(parts)
def insert_rows(wb,tgt,P,N):
    ws=wb[tgt]; mr=ws.max_row; mc=ws.max_column
    hts={i:ws.row_dimensions[i].height for i in range(P,mr+1)}
    ws.move_range(f'A{P}:{CL(mc)}{mr}',rows=N,cols=0,translate=False)
    for i,h in hts.items(): ws.row_dimensions[i+N].height=h
    for i in range(P,P+N): ws.row_dimensions[i].height=ws.row_dimensions[P-1].height
    for s in wb.worksheets:
        for row in s.iter_rows():
            for c in row:
                v=c.value
                if isinstance(v,str) and v.startswith('=') and (tgt in v or s.title==tgt):
                    nv=shift(v,s.title,tgt,P,N)
                    if nv!=v: c.value=nv
        for dv in s.data_validations.dataValidation:
            if dv.formula1 and tgt in dv.formula1: dv.formula1=shift(dv.formula1,s.title,tgt,P,N)
    if ws.auto_filter.ref:
        a,b=ws.auto_filter.ref.split(':'); col=re.match(r'[A-Z]+',b).group(0); r=int(b[len(col):])
        if r>=P-1: ws.auto_filter.ref=f'{a}:{col}{r+N}'
SING=re.compile(r"(?<![A-Za-z0-9_.:])((?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!)?(\$?[A-Z]{1,3}\$?)(\d+)(?![A-Za-z0-9_(!:])")
def row_formula(v,src,dst):
    if not(isinstance(v,str) and v.startswith('=')): return v
    parts=v.split('"')
    for q in range(0,len(parts),2):
        parts[q]=SING.sub(lambda m: f"{m.group(1) or ''}{m.group(2)}{dst}" if m.group(3)==str(src) else m.group(0),parts[q])
        parts[q]=re.sub(rf'(\$?[A-Z]{{1,3}}){src}:(\$?[A-Z]{{1,3}}){src}\b',lambda m:f'{m.group(1)}{dst}:{m.group(2)}{dst}',parts[q])
    return '"'.join(parts)
def regional_end(wb):
    ws=wb['Opp_Teams']
    for r in range(5,ws.max_row+1):
        if ws.cell(r,2).value=='Custom': return r-1
    raise RuntimeError('Opp_Teams: custom-opponent row not found')
# ---------------- add one event ----------------
def add_event(wb,name,teams,usage,log):
    N=len(teams)
    # Regional_Meta
    ws=wb['Regional_Meta']; last=4
    for r in range(5,ws.max_row+1):
        if ws.cell(r,1).value: last=r
    r=last+1; tpl=last
    for t in teams:
        for m in t['mons']:
            mv=m['mv']
            vals=[name,t['player'],t['place'],t['day'],m['dn'],m['sh'],m['pn'],m['cp'],mv[0][0],mv[0][1],mv[1][0],mv[1][1],mv[2][0] or None,mv[2][1] or None,m['key']]
            for j,v in enumerate(vals,1): c=ws.cell(r,j,v); cstyle(ws.cell(tpl,j),c)
            r+=1
    ws.auto_filter.ref=f'A4:W{r-1}'
    # Dracoviz_Usage
    du=wb['Dracoviz_Usage']; last=4
    for rr in range(5,du.max_row+1):
        if du.cell(rr,2).value: last=rr
    rr=last+1
    for sp,u,g in usage:
        du.cell(rr,1,f'=B{rr}&"|"&C{rr}'); du.cell(rr,2,name); du.cell(rr,3,sp); du.cell(rr,4,round(u,3)); du.cell(rr,5,round(g,3) if g is not None else None)
        for j in range(1,6): cstyle(du.cell(last,j),du.cell(rr,j))
        rr+=1
    du.auto_filter.ref=f'B4:E{rr-1}'
    cap=((rr+99)//100)*100
    ru=wb['Regional_Usage']
    for row in ru.iter_rows():
        for c in row:
            if isinstance(c.value,str) and 'Dracoviz_Usage!' in c.value:
                c.value=re.sub(r'(Dracoviz_Usage!\$[AE]\$5:\$[AE]\$)(\d+)',lambda m:m.group(1)+str(max(int(m.group(2)),cap)),c.value)
    # Regional_Usage: find event usage/GWR column pairs
    pairs=[]
    for col in range(1,ru.max_column+1):
        h=ru.cell(5,col).value
        if isinstance(h,str) and re.fullmatch(r'=([A-Z]+)3&" usage"',h):
            ev=ru.cell(3,col).value
            g=next((c2 for c2 in range(1,ru.max_column+1) if ru.cell(5,c2).value==f'GWR {ev}'),None)
            pairs.append((col,g))
    used=max(c for p in pairs for c in p)
    uc,gc=used+1,used+2; U,G=CL(uc),CL(gc); last_pair=pairs[-1]
    ru.cell(3,uc,name); cstyle(ru.cell(3,last_pair[0]),ru.cell(3,uc))
    ru.cell(4,uc,f'=SUMIFS(Regional_Meta!$Q$5:$Q$3004,Regional_Meta!$A$5:$A$3004,{U}$3)'); cstyle(ru.cell(4,last_pair[0]),ru.cell(4,uc))
    ru.cell(5,uc,f'={U}3&" usage"'); cstyle(ru.cell(5,last_pair[0]),ru.cell(5,uc))
    ru.cell(5,gc,f'GWR {name}'); cstyle(ru.cell(5,last_pair[1]),ru.cell(5,gc))
    pairs.append((uc,gc))
    existing=[ru.cell(i,1).value for i in range(6,401) if ru.cell(i,1).value]
    keymap={m['pn']:m['key'] for t in teams for m in t['mons']}
    new=sorted({m['pn'] for t in teams for m in t['mons']}-set(existing))
    rr=6+len(existing)
    for n in new:
        ru.cell(rr,1,n); ru.cell(rr,18,keymap[n]); cstyle(ru.cell(rr-1,1),ru.cell(rr,1)); cstyle(ru.cell(rr-1,18),ru.cell(rr,18)); rr+=1
    for i in range(6,401):
        ru.cell(i,uc,f'=IF($A{i}="","",COUNTIFS(Regional_Meta!$G$5:$G$3004,$A{i},Regional_Meta!$A$5:$A$3004,{U}$3)/{U}$4)')
        ru.cell(i,gc,f'=IF($R{i}="","",IFERROR(INDEX(Dracoviz_Usage!$E$5:$E${cap},MATCH({U}$3&"|"&$R{i},Dracoviz_Usage!$A$5:$A${cap},0)),""))')
        cstyle(ru.cell(i,last_pair[0]),ru.cell(i,uc)); cstyle(ru.cell(i,last_pair[1]),ru.cell(i,gc))
        num='+'.join(f'{CL(a)}{i}*{CL(a)}$4*N({CL(b)}{i})' for a,b in pairs); den='+'.join(f'{CL(a)}{i}*{CL(a)}$4*({CL(b)}{i}<>"")' for a,b in pairs)
        ru.cell(i,11,f'=IF($A{i}="","",IFERROR(({num})/({den}),""))')
    ru.column_dimensions[U].width=ru.column_dimensions[CL(last_pair[0])].width or 12; ru.column_dimensions[G].width=ru.column_dimensions[CL(last_pair[1])].width or 12
    ru.auto_filter.ref=f'A5:{G}{rr-1}'
    mt=wb['Meta_Threats']
    if isinstance(mt['I3'].value,str):
        mt['I3']=re.sub(r'COUNTA\(Regional_Usage![^)]*\)','COUNTA('+','.join(f'Regional_Usage!${CL(a)}$3' for a,_ in pairs)+')',mt['I3'].value)
    # Opp_Teams + Preview_Calc rows
    P=regional_end(wb)+1
    insert_rows(wb,'Opp_Teams',P,N); insert_rows(wb,'Preview_Calc',P,N)
    ot=wb['Opp_Teams']; mc=ot.max_column
    for k,t in enumerate(teams):
        r=P+k; vals=[f"{name} #{t['place']} {t['player']}",name,t['place'],t['player'],t['day']]
        for m in t['mons']:
            mv=m['mv']; vals+=[m['pn'],mv[0][1],mv[1][1],mv[2][1] or None,' / '.join(x[0] for x in mv if x[0]),mv[0][0],mv[1][0],mv[2][0] or None]
        for j,v in enumerate(vals,1): ot.cell(r,j,v)
        ot.cell(r,54,f'=COUNTIF($B$5:$B${P-1+N},B{r})'); ot.cell(r,55,f'="|"&F{r}&"|"&N{r}&"|"&V{r}&"|"&AD{r}&"|"&AL{r}&"|"&AT{r}&"|"')
        for j in range(1,mc+1): cstyle(ot.cell(P-1,j),ot.cell(r,j))
    pc=wb['Preview_Calc']; mc=pc.max_column; tmpl=[pc.cell(P-1,j).value for j in range(1,mc+1)]
    for k in range(N):
        r=P+k
        for j,v in enumerate(tmpl,1): pc.cell(r,j,row_formula(v,P-1,r)); cstyle(pc.cell(P-1,j),pc.cell(r,j))
    for r in range(5,pc.max_row+1):
        v=pc.cell(r,1).value
        if isinstance(v,str) and v.isdigit(): pc.cell(r,1,str(r-4))
        elif isinstance(v,int) and not isinstance(v,bool) and r<P+N+40: pc.cell(r,1,r-4)
    pf=wb['Partner_Finder']; endAK=None
    for row in pf.iter_rows():
        for c in row:
            if isinstance(c.value,str):
                m=re.search(r'\$AK\$8:\$AK\$(\d+)',c.value)
                if m: endAK=int(m.group(1))
    if endAK:
        tplc=pf.cell(endAK,37)
        for k in range(N):
            r=P+k; ar=r+3
            pf.cell(ar,37,f'=IF(ISNUMBER(SEARCH("|"&$B$4&"|",Opp_Teams!$BC${r})),Opp_Teams!$C${r}/Opp_Teams!$BB${r}+ROW()/10000000,9)'); cstyle(tplc,pf.cell(ar,37))
        for row in pf.iter_rows():
            for c in row:
                if isinstance(c.value,str) and f'$AK$8:$AK${endAK}' in c.value: c.value=c.value.replace(f'$AK$8:$AK${endAK}',f'$AK$8:$AK${endAK+N}')
    log(f'   added {name}: {N} teams, new species {len(new)}')
    return P,N
# ---------------- core sheets ----------------
NAVY='1F3864'; BLUE='2F5597'
TITLE=Font(b=True,sz=14,color=NAVY); NOTE=Font(sz=9,color='404040'); SEC=Font(b=True,sz=12,color='FFFFFF'); SECF=PatternFill('solid',fgColor=NAVY)
HDR=Font(b=True,sz=9,color='FFFFFF'); HDRF=PatternFill('solid',fgColor=BLUE); DAT=Font(sz=9); INP=Font(b=True,sz=11,color='0000FF'); INPF=PatternFill('solid',fgColor='FFFF00')
GREY=Font(sz=8,color='808080'); WR=Alignment(wrap_text=True,vertical='top'); TOP=Alignment(vertical='top'); BOX=Border(bottom=Side(style='thin',color='D9D9D9'))
LEGEND=(f"Letters follow the core's order, from the average of 0, 1 and 2 shields: W = clear win (rating {500+MARGIN}+), CW = close win (501–{500+MARGIN-1}), "
        f"CL = close loss ({500-MARGIN+1}–499), L = clear loss ({500-MARGIN} or less). A close loss is not a bad matchup — shields, baits or IVs can flip it. "
        "Answers avoid clear losses first and close losses second.")
def write_cores(wb,st,ans,main,cands,regend,log):
    lab=lambda c:' + '.join(c)
    def fmt(lst): return '\n'.join(f"{n} — {' '.join(ts)}" for n,ts,_,_ in lst[:3])
    def combo(x,core):
        if not x: return ''
        names,ts,cov=x; return ' + '.join(names)+f"  ({' '.join(ts)}; {cov}/{len(core)} clearly beaten)"
    def spec_of(i):
        s,n=ans.cand[i]; return main[n] if s=='meta' else cands[n]
    rows=[]
    for s in sorted(st['singles'],key=lambda s:-s['n'])+sorted(st['cores'],key=lambda s:(len(s['core']),-s['n'])):
        if not all(m in ans.mi for m in s['core']): continue
        ct=ans.counters(list(s['core']),None,None)
        w1=' | '.join(analyse.why(spec_of(ct['common'][0][3]),main[m],m) for m in s['core']) if ct['common'] else ''
        w2=' | '.join(analyse.why(spec_of(ct['anti'][0][3]),main[m],m) for m in s['core']) if ct['anti'] else ''
        rows.append((s,ct,w1,w2))
    for nm in ('Core_Data','Core_Finder'):
        if nm in wb.sheetnames: del wb[nm]
    cd=wb.create_sheet('Core_Data',wb.sheetnames.index('Opp_Teams')+1); cd.sheet_properties.tabColor='808080'
    H=['Key (helper)','Size','Core','Teams','Day 2 teams','Day 2 rate','Conversion','Events seen','Answers from the meta','Anti-meta answers','Counter duo','Counter trio','Why the #1 meta answer works','Why the #1 anti-meta answer works','','Finder sort (helper)']
    for j,h in enumerate(H,1): c=cd.cell(1,j,h); c.font=HDR; c.fill=HDRF; c.alignment=WR
    for i,(s,ct,w1,w2) in enumerate(rows,2):
        vals=['|'+'|'.join(s['core'])+'|',len(s['core']),lab(s['core']),s['n'],s['d'],s['rate'],s['conv'],s['events'],fmt(ct['common']),fmt(ct['anti']),combo(ct['duo'],s['core']),combo(ct['trio'],s['core']) if ct['trio'] else '',w1,w2]
        for j,v in enumerate(vals,1): c=cd.cell(i,j,v); c.font=DAT; c.alignment=WR if j>=9 else TOP
        cd.cell(i,6).number_format='0%'; cd.cell(i,7).number_format='0.00'
        cd.cell(i,16,f'=IF(AND(B{i}>1,ISNUMBER(SEARCH("|"&Core_Finder!$B$4&"|",A{i})),OR(Core_Finder!$B$5="Duos & trios",AND(Core_Finder!$B$5="Duos only",B{i}=2),AND(Core_Finder!$B$5="Trios only",B{i}=3))),IF(Core_Finder!$B$6="Teams",D{i}+F{i}/10,F{i}+D{i}/1000)+ROW()/10000000,-1)').font=GREY
    ND=len(rows)+1
    for j,w in enumerate([18,5,40,7,7,8,9,7,34,34,40,48,60,60,2,12],1): cd.column_dimensions[CL(j)].width=w
    cd.freeze_panes='D2'; cd.auto_filter.ref=f'A1:N{ND}'
    rng=lambda col:f'Core_Data!${col}$2:${col}${ND}'
    OPP=f'Opp_Teams!$BC$5:$BC${regend}'; OPPE=f'Opp_Teams!$E$5:$E${regend}'
    # Meta_Cores
    duomin,triomin=8,5; tab=None
    if 'Meta_Cores' in wb.sheetnames:
        old=wb['Meta_Cores']; tab=old.sheet_properties.tabColor
        if isinstance(old['G4'].value,(int,float)) and old['E4'].value and 'duos' in str(old['E4'].value): duomin=old['G4'].value
        elif isinstance(old['H4'].value,(int,float)): duomin=old['H4'].value
        if isinstance(old['J4'].value,(int,float)): triomin=old['J4'].value
        pos=wb.sheetnames.index('Meta_Cores'); del wb['Meta_Cores']
    else: pos=2
    ws=wb.create_sheet('Meta_Cores',pos)
    if tab: ws.sheet_properties.tabColor=tab
    NT=st['NT']
    ws['A1']='META CORES — the duos and trios top players built around, and what answers them'; ws['A1'].font=TITLE
    ws['A2']=(f"Counted live from the {NT} real regional teams (Opp_Teams). Teams = teams containing the whole core; Day 2 rate = share that reached Day 2; Conversion = Day 2 rate ÷ the average (above 1 = over-performs). "
      "Best-converting lists only rank cores on at least the yellow minimum number of teams. ANSWERS come from simulated battles (0, 1 and 2 shields) against each core member's most-used regional build, so typing, resisted moves, stats and energy all count. "
      +LEGEND+" 'Answers from the meta' = on 3+ regional teams; 'Anti-meta answers' = on fewer than 3 AND a net score of −20% or better vs the regional meta (wins minus losses, weighted by use). For moveset experiments use the Core Lab on the GL Core Scout site.")
    ws['A2'].font=NOTE; ws['A2'].alignment=WR; ws.merge_cells('A2:N2'); ws.row_dimensions[2].height=78
    ws['A4']='Average Day 2 rate:'; ws['A4'].font=Font(b=True,sz=9)
    ws['C4']=f'=COUNTIF({OPPE},"Day 2")/COUNTA({OPP})'; ws['C4'].number_format='0.0%'; ws['C4'].font=Font(b=True,sz=11)
    ws['E4']='Min teams — duos:'; ws['E4'].font=Font(b=True,sz=9); ws['G4']=duomin
    ws['I4']='Min teams — trios:'; ws['I4'].font=Font(b=True,sz=9); ws['J4']=triomin
    for a in ('G4','J4'): ws[a].font=INP; ws[a].fill=INPF; ws[a].alignment=Alignment(horizontal='center')
    duos=[s for s in st['cores'] if len(s['core'])==2]; trios=[s for s in st['cores'] if len(s['core'])==3]
    def helper(col0,lst,minref):
        cols=[CL(col0+k) for k in range(6)]; L,T,D2,R,K1,K2=cols
        for k,h in enumerate(['Candidate core','Teams','Day 2','Rate','Common key','Convert key']): c=ws.cell(6,col0+k,h); c.font=HDR; c.fill=HDRF
        for i,s in enumerate(lst,7):
            crit=','.join(f'{OPP},"*|{m}|*"' for m in s['core'])
            ws[f'{L}{i}']=lab(s['core']); ws[f'{T}{i}']=f'=COUNTIFS({crit})'; ws[f'{D2}{i}']=f'=COUNTIFS({crit},{OPPE},"Day 2")'
            ws[f'{R}{i}']=f'=IF({T}{i}=0,"",{D2}{i}/{T}{i})'; ws[f'{K1}{i}']=f'={T}{i}+N({R}{i})/10+ROW()/10000000'; ws[f'{K2}{i}']=f'=IF({T}{i}<{minref},-1,N({R}{i})+{T}{i}/1000+ROW()/10000000)'
            for c in cols: ws[f'{c}{i}'].font=GREY
        last=6+len(lst); return {k:f'${c}$7:${c}${last}' for k,c in zip(['L','T','D','R','K1','K2'],cols)}
    HD=helper(27,duos,'$G$4'); HT=helper(35,trios,'$J$4')
    HEAD=['Rank','Core','Teams','% of teams','Day 2 teams','Day 2 rate','Conversion','Events seen','Answers from the meta','Anti-meta answers','Counter duo','Counter trio','Why the #1 meta answer works','Why the #1 anti-meta answer works']
    r=[6]
    def section(title,Hh,key,trio):
        c=ws.cell(r[0],1,title); c.font=SEC
        for j in range(1,15): ws.cell(r[0],j).fill=SECF
        r[0]+=1
        for j,h in enumerate(HEAD,1): c=ws.cell(r[0],j,h); c.font=HDR; c.fill=HDRF; c.alignment=Alignment(wrap_text=True,vertical='center')
        r[0]+=1
        for k in range(1,51):
            rr=r[0]; P=f'$P{rr}'
            ws.cell(rr,16,f'=IF(LARGE({Hh[key]},{k})<0,"",MATCH(LARGE({Hh[key]},{k}),{Hh[key]},0))').font=GREY
            ws.cell(rr,1,k); ws.cell(rr,2,f'=IF({P}="","",INDEX({Hh["L"]},{P}))'); ws.cell(rr,3,f'=IF({P}="","",INDEX({Hh["T"]},{P}))')
            ws.cell(rr,4,f'=IF({P}="","",C{rr}/COUNTA({OPP}))'); ws.cell(rr,4).number_format='0.0%'
            ws.cell(rr,5,f'=IF({P}="","",INDEX({Hh["D"]},{P}))'); ws.cell(rr,6,f'=IF({P}="","",INDEX({Hh["R"]},{P}))'); ws.cell(rr,6).number_format='0%'
            ws.cell(rr,7,f'=IF({P}="","",IFERROR(F{rr}/$C$4,""))'); ws.cell(rr,7).number_format='0.00'
            for j,col in [(8,'H'),(9,'I'),(10,'J'),(11,'K'),(12,'L'),(13,'M'),(14,'N')]:
                if j==12 and not trio: continue
                ws.cell(rr,j,f'=IF($B{rr}="","",IFERROR(INDEX({rng(col)},MATCH($B{rr},{rng("C")},0))&"",""))')
            for j in range(1,15): c=ws.cell(rr,j); c.font=DAT; c.alignment=WR if j>=9 else TOP; c.border=BOX
            ws.row_dimensions[rr].height=40; r[0]+=1
        r[0]+=1
    section('MOST COMMON DUOS (top 50)',HD,'K1',False); section('BEST-CONVERTING DUOS (top 50, min teams in G4)',HD,'K2',False)
    section('MOST COMMON TRIO CORES (top 50)',HT,'K1',True); section('BEST-CONVERTING TRIO CORES (top 50, min teams in J4)',HT,'K2',True)
    vers,nbig=analyse.versatility(ans,st['cores'])
    c=ws.cell(r[0],1,'MOST VERSATILE ANSWERS — who avoids clear losses against the most cores'); c.font=SEC
    for j in range(1,15): ws.cell(r[0],j).fill=SECF
    r[0]+=1
    ws.cell(r[0],1,f"Checked against all {nbig} duo and trio cores used on 5+ teams, weighted by how many teams ran each. 'Covered' = no clear loss to any member (close losses allowed); 'Beaten' = a clear win against every member.").font=NOTE
    ws.merge_cells(start_row=r[0],start_column=1,end_row=r[0],end_column=14); ws.cell(r[0],1).alignment=WR; ws.row_dimensions[r[0]].height=28; r[0]+=1
    for j,h in enumerate(['Rank','Pokémon','Kind','Regional teams','Core teams covered','Core teams beaten outright'],1): c=ws.cell(r[0],j,h); c.font=HDR; c.fill=HDRF; c.alignment=Alignment(wrap_text=True,vertical='center')
    r[0]+=1
    for k,v in enumerate(vers[:50],1):
        for j,x in enumerate([k,v[0],v[1],st['use'].get(v[0],0),v[2],v[3]],1):
            c=ws.cell(r[0],j,x); c.font=DAT; c.border=BOX
            if j>=5: c.number_format='0%'
        r[0]+=1
    for j,w in enumerate([5,34,7,8,8,8,9,7,30,30,34,40,48,48,2,5],1): ws.column_dimensions[CL(j)].width=w
    ws.column_dimensions['P'].hidden=True
    for c in range(27,41): ws.column_dimensions[CL(c)].hidden=True
    ws.freeze_panes='C6'
    # Core_Finder
    cf=wb.create_sheet('Core_Finder',wb.sheetnames.index('Meta_Cores')+1)
    if tab: cf.sheet_properties.tabColor=tab
    cf['A1']='CORE FINDER — build around one Pokémon'; cf['A1'].font=TITLE
    cf['A2']=("Pick a Pokémon (yellow). Top: what beats it on its own. Below: every duo and trio core containing it that appeared on 3+ regional teams, with results and answers. "+LEGEND)
    cf['A2'].font=NOTE; cf['A2'].alignment=WR; cf.merge_cells('A2:N2'); cf.row_dimensions[2].height=52
    for a,l,v in [('A4','Pokémon:','Melmetal'),('A5','Show:','Duos & trios'),('A6','Sort by:','Teams')]:
        cf[a]=l; cf[a].font=Font(b=True,sz=10); b='B'+a[1:]; cf[b]=v; cf[b].font=INP; cf[b].fill=INPF
    for f,cell in [('Regional_Usage!$A$6:$A$400','B4'),('"Duos & trios,Duos only,Trios only"','B5'),('"Teams,Day 2 rate"','B6')]:
        dv=DataValidation(type='list',formula1=f); cf.add_data_validation(dv); dv.add(cell)
    cf['D4']='=IFERROR("On "&INDEX('+rng('D')+',MATCH("|"&$B$4&"|",'+rng('A')+',0))&" teams, "&TEXT(INDEX('+rng('F')+',MATCH("|"&$B$4&"|",'+rng('A')+',0)),"0%")&" reached Day 2 (average "&TEXT(Meta_Cores!$C$4,"0%")&").","Not in the regional data — pick a Pokémon used at the regionals.")'
    cf['D4'].font=Font(b=True,sz=10,color=NAVY)
    rr=8; c=cf.cell(rr,1,'WHAT BEATS IT ON ITS OWN'); c.font=SEC
    for j in range(1,15): cf.cell(rr,j).fill=SECF
    rr+=1
    for (j,h,col),(a,b) in zip([(1,'Answers from the meta','I'),(4,'Anti-meta answers','J'),(7,'Why the #1 meta answer works','M'),(11,'Why the #1 anti-meta answer works','N')],[(1,3),(4,6),(7,10),(11,14)]):
        c=cf.cell(rr,a,h); c.font=HDR
        for x in range(a,b+1): cf.cell(rr,x).fill=HDRF
        cf.merge_cells(start_row=rr,start_column=a,end_row=rr,end_column=b); cf.merge_cells(start_row=rr+1,start_column=a,end_row=rr+1,end_column=b)
        cc=cf.cell(rr+1,a,f'=IFERROR(INDEX({rng(col)},MATCH("|"&$B$4&"|",{rng("A")},0))&"","")'); cc.font=DAT; cc.alignment=WR
    cf.row_dimensions[rr+1].height=70; rr+=3
    c=cf.cell(rr,1,'CORES CONTAINING IT'); c.font=SEC
    for j in range(1,15): cf.cell(rr,j).fill=SECF
    rr+=1
    for j,h in enumerate(HEAD,1): c=cf.cell(rr,j,h); c.font=HDR; c.fill=HDRF; c.alignment=Alignment(wrap_text=True,vertical='center')
    rr+=1; PR=rng('P')
    for k in range(1,61):
        P=f'$P{rr}'
        cf.cell(rr,16,f'=IF(LARGE({PR},{k})<0,"",MATCH(LARGE({PR},{k}),{PR},0))').font=GREY
        cf.cell(rr,1,f'=IF({P}="","",{k})')
        for j,col in [(2,'C'),(3,'D'),(5,'E'),(6,'F'),(8,'H'),(9,'I'),(10,'J'),(11,'K'),(12,'L'),(13,'M'),(14,'N')]:
            cf.cell(rr,j,f'=IF({P}="","",INDEX({rng(col)},{P})&"")' if j>=9 else f'=IF({P}="","",INDEX({rng(col)},{P}))')
        cf.cell(rr,4,f'=IF({P}="","",C{rr}/COUNTA({OPP}))'); cf.cell(rr,4).number_format='0.0%'; cf.cell(rr,6).number_format='0%'
        cf.cell(rr,7,f'=IF({P}="","",IFERROR(F{rr}/Meta_Cores!$C$4,""))'); cf.cell(rr,7).number_format='0.00'
        for j in range(1,15): c=cf.cell(rr,j); c.font=DAT; c.alignment=WR if j>=9 else TOP; c.border=BOX
        cf.row_dimensions[rr].height=40; rr+=1
    for j,w in enumerate([5,34,7,8,8,8,9,7,30,30,34,40,48,48,2,5],1): cf.column_dimensions[CL(j)].width=w
    cf.column_dimensions['P'].hidden=True; cf.freeze_panes='A8'
    log(f'   core sheets: {len(rows)} cores and singles')
# ---------------- text touch-ups ----------------
def fix_text(wb,old_nt,new_nt,events):
    evs=', '.join(events[:-1])+' and '+events[-1]
    for ws in wb.worksheets:
        if ws.title.startswith(('Sim_','X_','Pokedex','Move','Slot','Regional_Meta','Opp_Teams','Preview_Calc')): continue
        for row in ws.iter_rows():
            for c in row:
                v=c.value
                if isinstance(v,str) and not v.startswith('='):
                    nv=v
                    if old_nt!=new_nt: nv=nv.replace(f'{old_nt} real teams',f'{new_nt} real teams').replace(f'all {old_nt} real',f'all {new_nt} real')
                    nv=re.sub(r'teams from [A-Z][a-z]+(, [A-Z][a-z]+)* and [A-Z][a-z]+',f'teams from {evs}',nv)
                    if nv!=v: c.value=nv
                elif isinstance(v,str) and v.startswith('=_xlfn._LONGTEXT('):
                    c.value=''.join(p.replace('""','"') for p in re.findall(r'"((?:[^"]|"")*)"',v[len('=_xlfn._LONGTEXT('):-1]))
# ---------------- Excel-only extensions (dropdowns, CF) ----------------
def sheet_files(z):
    wbx=z.read('xl/workbook.xml').decode('utf-8'); rels=z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    rm={re.search(r'Id="([^"]+)"',t).group(1):re.search(r'Target="([^"]+)"',t).group(1) for t in re.findall(r'<Relationship [^>]*>',rels)}
    out={}
    for t in re.findall(r'<sheet [^>]*>',wbx):
        n=re.search(r'name="([^"]+)"',t).group(1); i=re.search(r'r:id="([^"]+)"',t).group(1); tg=rm[i]
        out[n.replace('&amp;','&')]=tg[1:] if tg.startswith('/') else 'xl/'+tg
    return out
def reinject(src,dst,inserts):
    zs=zipfile.ZipFile(src); so=sheet_files(zs); exts={}
    for n,p in so.items():
        s=zs.read(p).decode('utf-8'); m=re.search(r'<extLst>.*</extLst>',s,re.S)
        if m:
            e=re.sub(r' xr\d?:uid="[^"]*"','',m.group(0))
            def fx(mm):
                t=mm.group(1)
                for tgt,P,N in inserts: t=shift(t,n,tgt,P,N)
                return f'<xm:f>{t}</xm:f>'
            exts[n]=re.sub(r'<xm:f>([^<]*)</xm:f>',fx,e)
    if not exts: return 0
    zd=zipfile.ZipFile(dst); do=sheet_files(zd); tmp=dst+'.tmp'
    with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zo:
        for it in zd.infolist():
            data=zd.read(it.filename)
            for n,e in exts.items():
                if do.get(n)==it.filename:
                    s=data.decode('utf-8')
                    if '<extLst>' not in s: data=s.replace('</worksheet>',e+'</worksheet>').encode('utf-8')
            zo.writestr(it,data)
    zd.close(); os.replace(tmp,dst); return len(exts)
# ---------------- entry point ----------------
def update(cfg,teams,usage,ans,st,main,cands,out_dir,log):
    src=os.path.join(cfg['inputs'],cfg['workbook']); wb=openpyxl.load_workbook(src)
    have=set()
    ws=wb['Regional_Meta']
    for r in ws.iter_rows(min_row=5,max_col=1,values_only=True):
        if r[0]: have.add(r[0])
    old_nt=wb['Opp_Teams'].cell(5,54).value; old_nt=regional_end(wb)-4
    inserts=[]
    for ev in cfg['events']:
        if ev['name'] in have: continue
        T=[t for t in teams if t['event']==ev['name']]
        P,N=add_event(wb,ev['name'],T,usage[ev['name']],log); inserts+= [('Opp_Teams',P,N),('Preview_Calc',P,N)]
    regend=regional_end(wb)
    write_cores(wb,st,ans,main,cands,regend,log)
    fix_text(wb,old_nt,st['NT'],[e['name'] for e in cfg['events']])
    wb.calculation.fullCalcOnLoad=True
    dst=os.path.join(out_dir,os.path.basename(cfg['workbook'])); wb.save(dst)
    n=reinject(src,dst,inserts); log(f'   saved {dst} ({n} sheets of Excel dropdowns/rules restored)')
